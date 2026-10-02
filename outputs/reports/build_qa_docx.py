from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(r'D:\Project Code\ANN-LSTM-Exam')
md_path = root / 'outputs' / 'reports' / 'QA_exam_answer.md'
docx_path = root / 'outputs' / 'reports' / 'QA_exam_answer.docx'
text = md_path.read_text(encoding='utf-8-sig')

# ==========================================
# 基础样式设置：正式报告风格，方便直接交作业
# ==========================================
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for style_name, size, color, before, after in [
    ('Heading 1', 16, '2E74B5', 16, 8),
    ('Heading 2', 13, '2E74B5', 12, 6),
    ('Heading 3', 12, '1F4D78', 8, 4),
]:
    st = styles[style_name]
    st.font.name = 'Calibri'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.line_spacing = 1.10

# 页脚页码（用 Word 字段）
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('第 ')
run.font.size = Pt(9)
field_begin = OxmlElement('w:fldChar')
field_begin.set(qn('w:fldCharType'), 'begin')
instr = OxmlElement('w:instrText')
instr.set(qn('xml:space'), 'preserve')
instr.text = 'PAGE'
field_end = OxmlElement('w:fldChar')
field_end.set(qn('w:fldCharType'), 'end')
footer._p.append(field_begin)
footer._p.append(instr)
footer._p.append(field_end)
run2 = footer.add_run(' 页')
run2.font.size = Pt(9)

# ==========================================
# Markdown 到 Word 的轻量转换函数
# ==========================================
def set_run_font(run, size=None, bold=None, italic=None, color=None, mono=False):
    run.font.name = 'Consolas' if mono else 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline_markdown(paragraph, raw_text):
    # 支持 **粗体** 和 `代码`，其余保持原文，避免引入复杂解析导致内容变形
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', raw_text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=10, mono=True, color='1F4D78')
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def clean_cell(value):
    value = value.strip()
    value = re.sub(r'\*\*([^*]+)\*\*', r'\1', value)
    value = re.sub(r'`([^`]+)`', r'\1', value)
    return value


def split_table_row(line):
    row = line.strip().strip('|')
    return [clean_cell(cell) for cell in row.split('|')]


def is_table_sep(line):
    s = line.strip()
    if not (s.startswith('|') and s.endswith('|')):
        return False
    body = s.strip('|').replace(' ', '')
    return bool(body) and all(ch in '-:|' for ch in body)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text_value, header=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (header or len(text_value) < 18) else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text_value)
    set_run_font(run, size=9, bold=header)


def add_table(rows):
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(col_count):
            cell = table.cell(i, j)
            set_cell_text(cell, row[j] if j < len(row) else '', header=(i == 0))
            if i == 0:
                set_cell_shading(cell, 'F2F4F7')
    for paragraph in doc.paragraphs[-1:]:
        paragraph.paragraph_format.space_after = Pt(6)


def add_horizontal_rule():
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'D9E2F3')
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

# ==========================================
# 正文生成
# ==========================================
lines = text.splitlines()
i = 0
in_code = False
code_lines = []

while i < len(lines):
    line = lines[i]

    if line.strip().startswith('```'):
        if not in_code:
            in_code = True
            code_lines = []
        else:
            in_code = False
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run('\n'.join(code_lines))
            set_run_font(run, size=9, mono=True, color='1F4D78')
        i += 1
        continue

    if in_code:
        code_lines.append(line)
        i += 1
        continue

    stripped = line.strip()
    if not stripped:
        i += 1
        continue

    # 表格块
    if stripped.startswith('|') and stripped.endswith('|') and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
        rows = [split_table_row(stripped)]
        i += 2
        while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
            rows.append(split_table_row(lines[i]))
            i += 1
        add_table(rows)
        continue

    if stripped == '---':
        add_horizontal_rule()
        i += 1
        continue

    if stripped.startswith('# '):
        p = doc.add_paragraph(style='Heading 1')
        add_inline_markdown(p, stripped[2:])
    elif stripped.startswith('## '):
        p = doc.add_paragraph(style='Heading 2')
        add_inline_markdown(p, stripped[3:])
    elif stripped.startswith('### '):
        p = doc.add_paragraph(style='Heading 3')
        add_inline_markdown(p, stripped[4:])
    elif stripped.startswith('> '):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(stripped[2:])
        set_run_font(run, italic=True, color='555555')
    elif re.match(r'^\d+\.\s+', stripped):
        p = doc.add_paragraph(style='List Number')
        add_inline_markdown(p, re.sub(r'^\d+\.\s+', '', stripped))
    elif stripped.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        add_inline_markdown(p, stripped[2:])
    else:
        p = doc.add_paragraph()
        add_inline_markdown(p, stripped)
    i += 1

# ==========================================
# 保存文档
# ==========================================
doc.core_properties.title = 'AI 增强型量化选股与回测系统 Q&A 实测报告'
doc.core_properties.subject = '基于 ANN-LSTM-Exam 实测输出生成'
doc.core_properties.author = 'Codex'
doc.save(docx_path)
print(docx_path)

